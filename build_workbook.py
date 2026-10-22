"""
Task 2a: results_tables.xlsx -- one sheet per table, T1-T7.
"""
import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

kpi_long = pd.read_csv("out/kpi_long.csv")

# T1: Overall, all indicators
t1 = kpi_long[kpi_long.stratifier == "Overall"][
    ["indicator", "percent", "ci_lo_per_100k", "ci_hi_per_100k", "rse_pct",
     "unweighted_n", "weighted_pop", "dof", "weight_var", "suppressed"]
].copy()
t1["ci_lo_per_100k"] = t1["ci_lo_per_100k"] / 1000
t1["ci_hi_per_100k"] = t1["ci_hi_per_100k"] / 1000
t1.columns = ["Indicator", "Percent", "CI lower (%)", "CI upper (%)", "RSE (%)",
              "Unweighted n", "Weighted population", "DoF", "Weight variable", "Suppressed"]

# T2: crude vs age-standardised
t2 = pd.read_csv("out/age_standardized.csv")
t2 = t2[t2.stratifier == "Overall"][
    ["indicator", "crude_percent", "crude_ci_lo", "crude_ci_hi", "crude_n",
     "std_percent", "std_ci_lo", "std_ci_hi", "std_n_age_groups_used",
     "std_n_age_groups_total", "crude_ci_half_width", "flag_diff_exceeds_half_width"]
].copy()
t2.columns = ["Indicator", "Crude %", "Crude CI lo", "Crude CI hi", "Crude n",
              "Std %", "Std CI lo", "Std CI hi", "Age groups used", "Age groups total",
              "Crude CI half-width", "Flag: |diff| > half-width"]

# T2b: crude vs standardised by sex/race/PIR (full detail, second sheet)
t2b = pd.read_csv("out/age_standardized.csv")
t2b = t2b[t2b.stratifier != "Overall"]

# T3: age x sex (pyramid feed)
t3 = kpi_long[kpi_long.stratifier == "AGE_BAND x SEX"][
    ["indicator", "stratum", "percent", "ci_lo_per_100k", "ci_hi_per_100k",
     "rse_pct", "unweighted_n", "suppressed"]
].copy()
t3["ci_lo_per_100k"] = t3["ci_lo_per_100k"] / 1000
t3["ci_hi_per_100k"] = t3["ci_hi_per_100k"] / 1000
t3.columns = ["Indicator", "Age band | Sex", "Percent", "CI lower (%)", "CI upper (%)",
              "RSE (%)", "Unweighted n", "Suppressed"]

# T4: concentration index, pooled + within-band
t4_pooled = pd.read_csv("out/concentration_index.csv")
t4_pooled.columns = ["Indicator", "Concentration index (pooled, adults 18+)"]
t4_band = pd.read_csv("out/t4_concentration_by_band.csv")
t4_band.columns = ["Indicator", "Age band", "Concentration index", "n in band"]

# T5: diagnosis gap
t5 = pd.read_csv("out/t5_diagnosis_gap.csv")
t5.columns = ["Measure", "Percent", "CI lower (%)", "CI upper (%)", "RSE (%)",
              "n", "DoF", "Weighted population", "Note"]

# T6: suppressed cells
t6 = pd.read_csv("out/t6_suppressed.csv")
t6.columns = ["Indicator", "Stratifier", "Stratum", "Unweighted n", "RSE (%)", "Reason"]

# T7: validation vs published
t7 = pd.read_csv("out/t7_validation.csv")
t7.columns = ["Indicator", "Mine (crude %)", "Mine (standardized %)", "Published",
              "Published age range", "Source", "URL", "Definitional differences"]

with pd.ExcelWriter("results_tables.xlsx", engine="openpyxl") as writer:
    t1.to_excel(writer, sheet_name="T1 Overall", index=False)
    t2.to_excel(writer, sheet_name="T2 Crude vs Standardized", index=False)
    t2b.to_excel(writer, sheet_name="T2b Std by subgroup", index=False)
    t3.to_excel(writer, sheet_name="T3 Age x Sex", index=False)
    t4_pooled.to_excel(writer, sheet_name="T4 Concentration Index", index=False)
    t4_band.to_excel(writer, sheet_name="T4 Concentration Index", index=False,
                      startrow=len(t4_pooled) + 3)
    t5.to_excel(writer, sheet_name="T5 Diagnosis Gap", index=False)
    t6.to_excel(writer, sheet_name="T6 Suppressed Cells", index=False)
    t7.to_excel(writer, sheet_name="T7 Validation", index=False)

    wb = writer.book
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=1):
            for cell in row:
                if cell.value is not None:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        for col_cells in ws.columns:
            length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
            col_letter = get_column_letter(col_cells[0].column)
            ws.column_dimensions[col_letter].width = min(max(length + 2, 10), 60)

    # second header row for the T4 sheet's second table
    ws4 = wb["T4 Concentration Index"]
    start = len(t4_pooled) + 4
    for cell in ws4[start]:
        if cell.value is not None:
            cell.fill = header_fill
            cell.font = header_font

print("Wrote results_tables.xlsx")
